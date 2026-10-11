"""build: rows.json -> SpawnPrawn_Copy.xlsx.   read: edited xlsx -> edits.json (only rows you changed)."""
import json, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEAD = ['Ref (do not edit)', 'Item', 'Field', 'Current text', 'NEW TEXT (type here)', 'Where', 'Same text elsewhere', 'Careful']
WID = [16, 26, 14, 60, 60, 16, 24, 28]

def build(rows_json, out):
    rows = json.load(open(rows_json)); wb = Workbook(); ws0 = wb.active; ws0.title = 'READ ME'
    tabs = []
    for r in rows:
        if r['tab'] not in tabs: tabs.append(r['tab'])
    hf = PatternFill('solid', fgColor='222222'); ef = PatternFill('solid', fgColor='FFF8D6'); grey = Font(color='888888', size=9)
    wrap = Alignment(wrap_text=True, vertical='top')
    meta = wb.create_sheet('_map'); meta.sheet_state = 'hidden'; meta.append(['ref', 'file', 'n', 'old', 'kind'])
    for t in tabs:
        ws = wb.create_sheet(t); ws.append(HEAD)
        for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = hf; c.alignment = Alignment(vertical='center')
        lastgroup = None
        for r in [x for x in rows if x['tab'] == t]:
            g = r['group'] or r['file']
            if g != lastgroup:                      # section band: file and table name, so you can jump about
                ws.append([None, f"{r['file']}  ›  {g}"]); b = ws.max_row
                for c in ws[b][:8]: c.fill = PatternFill('solid', fgColor='D9E2F3'); c.font = Font(bold=True)
                lastgroup = g
            ws.append([r['id'], r['item'], r['field'], r['text'], None, f"{r['file']}:{r['line']}", r['repeats'], r['note']])
            i = ws.max_row; ws.cell(i, 1).font = grey; ws.cell(i, 6).font = grey
            ws.cell(i, 5).fill = ef
            if r['note']: ws.cell(i, 8).font = Font(color='B45309')
            for c in range(1, 9): ws.cell(i, c).alignment = wrap
            meta.append([r['id'], r['file'], r['n'], r['text'], r['kind']])
        for i, w in enumerate(WID, 1): ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'C2'; ws.auto_filter.ref = f"A1:H{ws.max_row}"
    ws0.column_dimensions['A'].width = 32; ws0.column_dimensions['B'].width = 14; ws0.column_dimensions['C'].width = 90
    L = [('SPAWN PRAWN: ALL GAME COPY', '', ''), ('', '', ''),
         ('How to edit', '', 'Type your replacement in the yellow NEW TEXT column. Leave it empty to keep the current text. To make a line empty, type [blank].'),
         ('', '', 'Do not change the grey Ref column, and do not insert or reorder columns. Rows can be filtered and sorted freely.'),
         ('', '', 'Keep anything in ${...} or <...> exactly as it is (those are code). Rows with them are marked in the Careful column.'),
         ('', '', 'Names (weapons, enemies) are sometimes used elsewhere; the "Same text elsewhere" column shows where. Tell Claude if you rename something and it will check the matching lines.'),
         ('', '', 'To apply: save the file, share it with Claude, and say "apply the copy sheet".'),
         ('', '', ''), ('Tab', 'Rows', 'What is on it')]
    desc = {'1 Menus & Screens': 'Title, sample/sequence select, settings, HUD, summary and end screens, Gene Bank, Lab',
            '2 Tutorials & Hints': 'Tutorial cards, first-sighting enemy cards, status cards, glossary',
            '3 Sequences': 'The playable sequences: profiles, abilities, evolutions, unlocks, reveal stories, Born-again traits',
            '4 Weapons': 'Weapons, signature upgrades, pairings, combos, toys, modifiers',
            '5 Feats': 'Feats (spells) and their forks',
            '6 Cards & Upgrades': 'Passives, perks, curses, mutations, boons, junk DNA, power-ups, relics, dyes, stain grants, loot cards',
            '7 Damage Types': 'Damage types, chemical reactions, synergies, quirks, twists',
            '8 Enemies & Bosses': 'Enemies, bosses, their abilities and intros',
            '9 Rivals & Egg': 'Rival champions, their relics and quotes, the egg',
            '10 Modes & Levels': 'Samples, Petri Dish waves, wave conditions, campaign levels, run events, difficulty',
            '11 Narrator & Awards': 'Lab tech narrator lines, kill lines, achievements, sponsors, system messages',
            '12 Other': 'Anything that did not fit elsewhere'}
    L += [(t, sum(1 for x in rows if x['tab'] == t), desc.get(t, '')) for t in tabs]
    for a, b, c in L: ws0.append([a, b, c])
    for row in ws0.iter_rows():
        for c in row: c.alignment = wrap
    ws0['A1'].font = Font(bold=True, size=16); ws0['A3'].font = Font(bold=True)
    for c in ws0[9]: c.font = Font(bold=True)
    for i in range(10, 10 + len(tabs)):
        ws0.cell(i, 1).hyperlink = f"#'{ws0.cell(i, 1).value}'!A1"; ws0.cell(i, 1).font = Font(color='1D4ED8', underline='single')
    wb.move_sheet('_map', offset=len(wb.sheetnames) - 1 - wb.sheetnames.index('_map'))
    wb.save(out); print('saved', out, len(rows), 'strings')

def read(xlsx, out):
    wb = load_workbook(xlsx, data_only=True); m = {}
    for r in wb['_map'].iter_rows(min_row=2, values_only=True): m[r[0]] = r
    edits = []
    for ws in wb.worksheets:
        if ws.title in ('READ ME', '_map'): continue
        for r in ws.iter_rows(min_row=2, values_only=True):
            ref = r[0]
            if not ref or ref not in m: continue
            _, file, n, old, kind = m[ref]
            cur, new = r[3], r[4]
            if new is not None and str(new).strip() != '':
                new = '' if str(new).strip() == '[blank]' else str(new)
            elif cur is not None and str(cur) != old: new = str(cur)       # edited in Current text by mistake: still honour it
            else: continue
            if new != old: edits.append({'id': ref, 'file': file, 'n': n, 'old': old, 'new': new})
    json.dump(edits, open(out, 'w')); print(len(edits), 'changed rows')

if __name__ == '__main__':
    {'build': build, 'read': read}[sys.argv[1]](sys.argv[2], sys.argv[3])
